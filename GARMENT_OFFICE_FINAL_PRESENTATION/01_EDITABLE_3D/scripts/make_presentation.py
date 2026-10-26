"""Client presentation (18 slides, 16:9) from one content spec ->
03_PRESENTATION/Garment_Office_Design_Pitch.html (PDF via export_presentation.js) and Garment_Office_Design_Pitch.pptx (editable).
Uses only renders / drawings produced from this project's model.
python3 make_presentation.py PROJECT_ROOT
Images: 02_RENDERED_IMAGES/final/<CAM>.jpg if present, else 02_RENDERED_IMAGES/previews/<CAM>.jpg (caption says 'preview render').
"""
import html
import os
import shutil
import sys

ROOT = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.getcwd()
sys.path.insert(0, os.path.join(ROOT, "01_EDITABLE_3D", "scripts"))
import office_layout as O  # noqa: E402
import survey as S  # noqa: E402

OUT = os.path.join(ROOT, "03_PRESENTATION")
ASSET = os.path.join(OUT, "assets")
os.makedirs(ASSET, exist_ok=True)
FINAL = os.path.join(ROOT, "02_RENDERED_IMAGES", "final")
PREV = os.path.join(ROOT, "02_RENDERED_IMAGES", "previews")
PLANS = os.path.join(ROOT, "02_RENDERED_IMAGES", "plans")
BRAND = os.path.join(ROOT, "01_EDITABLE_3D", "assets", "brand")
USED = {}
E = html.escape


def img(cam):
    for folder, tag in ((FINAL, ""), (PREV, " (preview render)")):
        p = os.path.join(folder, cam + ".jpg")
        if os.path.exists(p):
            shutil.copyfile(p, os.path.join(ASSET, cam + ".jpg"))
            USED[cam] = "final" if not tag else "preview"
            return "assets/" + cam + ".jpg", tag
    USED[cam] = "missing"
    return None, " (not rendered)"


def asset(path, name=None):
    if path and os.path.exists(path):
        name = name or os.path.basename(path)
        shutil.copyfile(path, os.path.join(ASSET, name))
        return "assets/" + name
    return None


for f in ("logo_zara_fashion_profile.png", "logo_stage2_ltd.png", "logo_stage2_unit1.png", "tagline_quality_compliance_commitment.png"):
    asset(os.path.join(BRAND, f))
area = {z[1]: S.poly_area_sft(z[3]) for z in O.ZONES}
TOTAL_FIT = sum(area.values())
GROUP_NAMES = {"Reception": "Visitor & reception", "Experience_Center": "Experience centre", "Meeting_Rooms": "Meeting", "Executive_Area": "Executive",
               "Workstations": "Administration & workspaces", "Design_Development": "Design & development", "Support_Rooms": "Staff support & circulation",
               "Prayer_Rooms": "Prayer", "Washrooms": "Washrooms"}
groups = {}
for coll, zid, label, pts, fin, ceil in O.ZONES:
    groups[GROUP_NAMES[coll]] = groups.get(GROUP_NAMES[coll], 0) + area[zid]

# ------------------------------------------------------------------------------------------ content spec
SPEC = [
    dict(type="cover", image=("CAM_04_Experience_Centre_A", "Experience centre - Blender Cycles render of the project model"),
         eyebrow="Interior design proposal  |  Ground floor",
         title="Premium International Fashion Corporate Headquarters & Design Development Centre",
         sub=["Stage II Knitting Apparel Ltd.  ·  Stage II Knitting Apparel Unit-1  ·  Zara Fashion", "Holding 329/2, Ward 15, Vogra, Bason, Gazipur"]),
    dict(type="room", eyebrow="Project vision", title="One floor that tells the story from yarn to buyer-ready collection",
         text="The ground floor becomes the face of the company: a calm, premium headquarters where international buyers meet the brand, management "
              "works in comfort and the Global Design Studio turns trends into production-ready products - with no production on this floor.",
         bullets=["Visitor journey: lobby - reception - experience centre - boardroom", "Executive suite and administration beside the visitor zone",
                  "Design-development wing: studio, CAD, fabric library, 3D lab, sample room", "Staff cafe, lounge / training, prayer rooms and washrooms"],
         images=[("CAM_02_Reception_B", "Reception and visitor waiting")], layout="1"),
    dict(type="cards", eyebrow="Design philosophy", title="Six principles behind every decision", numbered=True,
         cards=[("Modernity", "Clean planes, slim black-aluminium glass, an exposed and painted structure - an honest contemporary shell."),
                ("Functionality", "Clear zoning: visitors, executives, administration, design development and staff support each have their place."),
                ("Collaboration", "Team tables, a buyer presentation table, a sample-review counter and glass views between departments."),
                ("Creativity", "Mood-board wall, fabric & trim library, 3D printing lab and display bays keep products and ideas visible."),
                ("Employee comfort", "Ergonomic 1400 mm desks, acoustic felt, 4000 K task light, staff cafe, lounge and prayer rooms."),
                ("Professional image", "The client's own logos and green, terrazzo, oak and walnut: a confident international brand presence.")]),
    dict(type="plan", plan="ID-01_furnished_plan.png", eyebrow="Floor plan", title="Furnished ground-floor layout",
         text=f"Overall 92'-4\" front x 96'-6\" (west) / 151'-3\" (east) outside dimensions, from the client's hand-measured sketches. "
              f"Fit-out zones total about {TOTAL_FIT:,.0f} sft net.",
         table=[("Key dimension", "Value", "Status"), ("Rooms 1 / 3 / 4 widths", "25'-0\" / 21'-8\" / 43'-2\"", "measured"),
                ("Middle block depth", "73'-5\"", "derived"), ("Rooms 5 + 6 width", "41'-7\"", "measured"),
                ("Room 8 length / width", "128'-2\" / 37'-6\"", "measured / assumed"), ("Floor to deck", "9'-6\"", "photo estimate")],
         note="Drawing set (A3): ID-01 furnished plan, ID-02 zoning, ID-03 ceiling + lighting, ID-04 power, data + fire. "
              "Geometry check: 75 of 75 model measurements match the survey values."),
    dict(type="plan", plan="ID-02_zoning_plan.png", eyebrow="Space zoning", title="Visitors in front, work in the middle, creation and care at the back",
         text="Red route: visitors. Green route: staff. Brown: executive link to the boardroom. Visitors never pass through the offices.",
         table=[("Zone", "Net area", "")] + [(k, f"{v:,.0f} sft", "") for k, v in groups.items()], note=""),
    dict(type="full", image=("CAM_ISO_SE", "Overall isometric cut-away of the furnished floor - Blender Cycles render"), eyebrow="Overall 3D view",
         title="The complete ground floor, furnished",
         text="Editable Blender model: every wall, partition, door, light and piece of furniture is a separate named object in named collections."),
    dict(type="room", eyebrow="Reception & entrance", title="A calm, branded welcome",
         text="Guests step down from the gate into a terrazzo lobby with a security desk and the group's three logos on an oak brand wall, then into a warm reception.",
         bullets=["Fluted white-oak counter, terrazzo ledge, accessible 760 mm section", "Backlit logo plate on a fluted oak wall",
                  "Waiting lounge under an oak slat raft and a brass halo", "Glass between lobby and reception"],
         images=[("CAM_01_Reception_A", "Reception"), ("CAM_02_Reception_B", "Reception - waiting"), ("CAM_03_Entrance_Lobby", "Entrance lobby - brand wall")]),
    dict(type="room", eyebrow="Experience centre", title="Where buyers meet the product",
         text="At the centre of the floor, between reception and boardroom, the experience centre tells the company story and shows the collection.",
         bullets=["Brand-story wall with the client's logos and tagline", "Four garment bays: deep green backs, brass rails, LED shelves",
                  "Mannequin plinth and 'yarn to garment' touch table", "Track lights fixed to the existing beams", "Glass views into boardroom and fabric library"],
         images=[("CAM_04_Experience_Centre_A", "Experience centre"), ("CAM_05_Experience_Centre_B", "Experience centre - story wall"), ("CAM_06_Experience_Centre_Detail", "Display detail")]),
    dict(type="room", eyebrow="Executive office", title="Managing Director's office",
         text="A quiet, warm executive suite next to the reception, reached through the PA and executive waiting area and a glass executive corridor.",
         bullets=["2.1 m smoked-walnut desk, cognac leather chair", "Walnut feature wall with the company logo", "Lounge corner for informal meetings",
                  "Oak herringbone floor, oak slat raft and halo light", "Direct link from the executive suite to the boardroom"],
         images=[("CAM_07_MD_Office_A", "MD office"), ("CAM_08_MD_Office_B", "MD office - lounge corner"), ("CAM_17_Executive_Reception", "Executive reception (PA)")]),
    dict(type="room", eyebrow="Boardroom", title="A boardroom for international buyers",
         text="Fifteen seats around a 6.1 m walnut table, between the reception, the executive suite and the experience centre.",
         bullets=["86 in screen on an oak slat media wall", "Green felt acoustic rafts set between the existing beams", "Walnut acoustic slat wall, herringbone oak floor",
                  "Power/data trough in the table; privacy curtains recommended on the glass"],
         images=[("CAM_09_Boardroom_A", "Boardroom"), ("CAM_10_Boardroom_B", "Boardroom - towards the experience centre")], layout="2"),
    dict(type="room", eyebrow="Main office", title="Workplaces that work",
         text="Merchandising, commercial & accounts and HR & admin share one desk family and one light and acoustic language.",
         bullets=["48 dedicated workstations + 6 guest seats (about 50 % growth)", "1400 x 700 mm bench desks, green felt screens, task chairs",
                  "Team table and sample-review counter for merchandisers", "Felt baffles between the beams, 4000 K linear light"],
         images=[("CAM_11_Merchandising_A", "Merchandising open office"), ("CAM_12_Merchandising_B", "Merchandising - team table"), ("CAM_16_Commercial_Accounts", "Commercial & accounts")]),
    dict(type="room", eyebrow="Design studio & CAD", title="The Global Design Studio",
         text="Zones from the master plan (about 1,400 sft): creative desks, trend wall, collection presentation, technical design, fabric library and storage.",
         bullets=["Designer desks at the east windows", "7.4 m mood-board wall for trend research", "Buyer presentation table with 86 in screen",
                  "CAD room with plotter; fabric & trim library with D65 booth", "3D printing lab and sample room as office-grade support"],
         images=[("CAM_13_Design_Studio_A", "Global Design Studio"), ("CAM_14_Design_Studio_B", "Design studio - workstations + mood board"), ("CAM_15_CAD_Room", "CAD / technical design")]),
    dict(type="grid", eyebrow="Support spaces", title="Care for the people who make the product",
         images=[("CAM_20_Staff_Cafe", "Staff cafe + pantry"), ("CAM_21_Lounge_Training", "Lounge / training (multipurpose)"), ("CAM_22_Prayer_Male", "Prayer room (male)"),
                 ("CAM_23_Sample_Room", "Sample development"), ("CAM_25_Fabric_Library", "Fabric & trim library"), ("CAM_28_Washroom_Male", "Staff washroom")]),
    dict(type="materials", eyebrow="Materials & finishes", title="A warm, durable palette",
         table=[("Code", "Finish", "Where"), ("TZ-01", "Terrazzo-look porcelain, warm beige", "visitor floors"), ("WD-01", "Oak herringbone", "boardroom, executive"),
                ("CP-01", "Charcoal carpet tile", "offices"), ("WO-01", "White-oak slats on acoustic felt", "feature walls"), ("WO-02", "Smoked walnut veneer", "executive furniture"),
                ("PA-02", "Deep green from the logo", "accents"), ("GL-01", "12 mm glass, black aluminium", "partitions"), ("MT-02", "Brushed brass", "details")],
         note="Specification concepts; final products by sample approval. Full schedule in 04_DOCUMENTATION.",
         image=("material_board.jpg", "Material board - the PBR materials of the model, rendered in Blender")),
    dict(type="plan", plan="ID-03_ceiling_lighting_plan.png", eyebrow="Lighting & ceiling", title="The beams stay - and do the work",
         bullets=["Existing steel deck and beams kept exposed, painted warm white", "Experience centre: track lights fixed to the beams",
                  "Boardroom: green felt rafts between the beams; merchandising: felt baffles", "Reception and MD office: oak slat rafts below beam level",
                  "4000 K direct/indirect light over desks; 3000-3500 K in visitor and executive areas", "Accent spots on products and brand walls; halo pendants"],
         image=("CAM_09_Boardroom_A", "Felt rafts between the beams - boardroom")),
    dict(type="cards", eyebrow="Design benefits", title="What the company gains", numbered=False,
         cards=[("Operational efficiency", "Departments placed by workflow: merchandising beside the experience centre and sample room, HR at the staff entrance, the studio next to CAD, library and samples."),
                ("Professional image", "A visitor sequence that presents the brand, the product and the management in one walk."),
                ("Employee comfort", "Ergonomic desks, acoustic control, daylight in the studio, cafe, lounge and prayer rooms."),
                ("Flexibility", "Bench desks, a multipurpose lounge / training room, glass partitions and about 50 % seat growth."),
                ("Maintainability", "Hard-wearing terrazzo and porcelain, replaceable carpet tiles, exposed services that stay accessible.")],
         image=("CAM_13_Design_Studio_A", "Global Design Studio")),
    dict(type="table", eyebrow="Implementation considerations", title="To confirm before construction",
         table=[("Item", "What has to be confirmed"),
                ("Site measurements", "Floor-to-deck height, room 8 width, column positions and beam layout (currently estimated)."),
                ("Structure", "Lintels for the new openings D17, D20, D22 and the lobby-reception glass - structural engineer."),
                ("Fire & escape", "Ground-floor exits of the NE / SE stairs; without them the walk to the main gate is about 220 ft from the lounge."),
                ("MEP", "VRF cooling, fresh air for windowless rooms, exhausts (WCs, pantry, 3D lab), power/data, UPS, DB location."),
                ("Accessibility", "Ramp or lift at the two entrance steps; one accessible WC (door at least 36 in)."),
                ("Approvals & branding", "Fire / building approvals as applicable; trademark check for 'Global Design Studio'; signage from original vector logos.")],
         note="No budget or programme was supplied, so none is stated. Details: 04_DOCUMENTATION/technical_coordination_notes.md."),
    dict(type="hero", image=("CAM_05_Experience_Centre_B", "Experience centre towards the brand-story wall - Blender Cycles render"), eyebrow="Closing",
         title="Quality, compliance and commitment - built into the space.",
         text="A ground floor that welcomes buyers, supports management and gives the design team room to create - ready for detailed design "
              "once the open site measurements are confirmed."),
]

# ------------------------------------------------------------------------------------------ HTML renderer
CSS = """
@page { size: 1920px 1080px; margin: 0 }
* { box-sizing: border-box }
html, body { margin: 0; padding: 0; background: #2b2b2b }
body { font-family: 'Inter', 'Helvetica Neue', Arial, sans-serif; color: #1E1E1C }
.slide { width: 1920px; height: 1080px; position: relative; overflow: hidden; background: #F6F3EE; page-break-after: always; break-after: page }
.eyebrow { font-size: 19px; letter-spacing: .22em; text-transform: uppercase; color: #26673D; font-weight: 600 }
h1 { font-family: 'Inter Display','Inter',sans-serif; font-size: 70px; line-height: 1.04; margin: 14px 0 0; font-weight: 600; letter-spacing: -.015em }
h2 { font-family: 'Inter Display','Inter',sans-serif; font-size: 46px; line-height: 1.1; margin: 12px 0 0; font-weight: 600; letter-spacing: -.01em }
p, li { font-size: 23px; line-height: 1.5; color: #3A3935 }
ul { padding-left: 24px; margin: 16px 0 0 } li { margin: 0 0 9px } li::marker { color: #26673D }
.foot { position: absolute; left: 80px; right: 80px; bottom: 30px; display: flex; justify-content: space-between; font-size: 15px; color: #8A867E; letter-spacing: .04em }
.foot b { color: #26673D; font-weight: 600 }
.dark .foot { color: #A7B3AB } .dark .foot b { color: #9FD3B0 }
.bar { position: absolute; left: 80px; top: 70px; width: 64px; height: 4px; background: #26673D }
.imgbox { position: absolute; overflow: hidden; background: #d9d6d0 } .imgbox img { width: 100%; height: 100%; object-fit: cover; display: block }
.cap { position: absolute; left: 0; bottom: 0; padding: 9px 16px; background: rgba(20,20,18,.62); color: #F3F0EA; font-size: 15px; letter-spacing: .03em }
.card { position: absolute; background: #fff; border-radius: 6px; padding: 28px 28px 22px }
.card h3 { font-size: 29px; margin: 10px 0 10px; font-weight: 600 } .card p { font-size: 21px; margin: 0; line-height: 1.45 }
.num { font-family: 'Inter Display','Inter'; font-size: 50px; font-weight: 600; color: #26673D; line-height: 1 }
table { border-collapse: collapse; width: 100% } td, th { text-align: left; padding: 10px 12px; border-bottom: 1px solid #DDD7CB; font-size: 20px; vertical-align: top }
th { font-size: 14px; letter-spacing: .14em; text-transform: uppercase; color: #6B675F; font-weight: 600 }
.logo-plate { position: absolute; background: #fff; border-radius: 4px; display: flex; align-items: center; gap: 26px; padding: 18px 26px }
"""


def h_img(spec_img, x, y, w, h, fit="cover"):
    name, caption = spec_img
    if name.endswith(".jpg") or name.endswith(".png"):
        svg = os.path.join(PLANS, name[:-4] + ".svg")
        src = (asset(svg) if os.path.exists(svg) else None) or asset(os.path.join(ROOT, "02_RENDERED_IMAGES", name)) or asset(os.path.join(PLANS, name))
        asset(os.path.join(PLANS, name))            # keep the PNG too (used by the PPTX)
        tag = ""
    else:
        src, tag = img(name)
    if not src:
        return f'<div class="imgbox" style="left:{x}px;top:{y}px;width:{w}px;height:{h}px"></div>'
    cap = f'<div class="cap">{E(caption + tag)}</div>' if caption else ""
    return f'<div class="imgbox" style="left:{x}px;top:{y}px;width:{w}px;height:{h}px"><img src="{src}" style="object-fit:{fit}">{cap}</div>'


def h_table(rows, widths=None):
    head = "".join(f"<th>{E(c)}</th>" for c in rows[0])
    body = "".join("<tr>" + "".join(f"<td>{'<b>' + E(c) + '</b>' if i == 0 else E(c)}</td>" for i, c in enumerate(r)) + "</tr>" for r in rows[1:])
    return f"<table><tr>{head}</tr>{body}</table>"


def h_slide(n, s):
    t, inner, dark = s["type"], "", False
    if t == "cover":
        dark = True
        inner = (h_img(s["image"], 0, 0, 1920, 1080)
                 + '<div style="position:absolute;inset:0;background:linear-gradient(90deg,rgba(15,30,21,.93) 0%,rgba(15,30,21,.8) 38%,rgba(15,30,21,0) 68%)"></div>'
                 + '<div class="logo-plate" style="left:80px;top:80px"><img src="assets/logo_zara_fashion_profile.png" style="height:62px">'
                   '<img src="assets/logo_stage2_ltd.png" style="height:50px"><img src="assets/logo_stage2_unit1.png" style="height:50px"></div>'
                 + f'<div style="position:absolute;left:80px;top:320px;width:930px;color:#F3F0EA"><div class="eyebrow" style="color:#9FD3B0">{E(s["eyebrow"])}</div>'
                   f'<h1 style="color:#fff;font-size:72px">{E(s["title"])}</h1>'
                   + "".join(f'<p style="color:#D9E2DC;font-size:25px;margin:{26 if i == 0 else 6}px 0 0">{E(l)}</p>' for i, l in enumerate(s["sub"]))
                   + '<img src="assets/tagline_quality_compliance_commitment.png" style="height:40px;margin-top:34px;background:#fff;padding:10px 18px;border-radius:3px"></div>')
    elif t == "room":
        inner = (f'<div class="bar"></div><div style="position:absolute;left:80px;top:104px;width:560px"><div class="eyebrow">{E(s["eyebrow"])}</div>'
                 f'<h2>{E(s["title"])}</h2><p>{E(s["text"])}</p><ul>' + "".join(f"<li>{E(b)}</li>" for b in s["bullets"]) + "</ul></div>")
        ims = s["images"]
        lay = s.get("layout", "1+2")
        if lay == "1":
            inner += h_img(ims[0], 700, 104, 1140, 896)
        elif lay == "2":
            inner += h_img(ims[0], 700, 104, 1140, 440) + h_img(ims[1], 700, 560, 1140, 440)
        else:
            inner += h_img(ims[0], 700, 104, 1140, 560) + h_img(ims[1], 700, 680, 562, 320) + h_img(ims[2], 1278, 680, 562, 320)
    elif t == "cards":
        inner = f'<div class="bar"></div><div style="position:absolute;left:80px;top:104px;width:1600px"><div class="eyebrow">{E(s["eyebrow"])}</div><h2>{E(s["title"])}</h2></div>'
        for i, (ct, cd) in enumerate(s["cards"]):
            num = f'<div class="num">{i + 1:02d}</div>' if s.get("numbered") else ""
            inner += (f'<div class="card" style="left:{80 + (i % 3) * 600}px;top:{320 + (i // 3) * 300}px;width:560px;height:{270 if s.get("numbered") else 250}px">'
                      f'{num}<h3>{E(ct)}</h3><p>{E(cd)}</p></div>')
        if s.get("image"):
            inner += h_img(s["image"], 1280, 620, 560, 250)
    elif t == "plan":
        inner = h_img((s["plan"], ""), 80, 56, 680, 962, fit="contain")
        inner += f'<div style="position:absolute;left:830px;top:104px;width:1010px"><div class="eyebrow">{E(s["eyebrow"])}</div><h2>{E(s["title"])}</h2>'
        if s.get("text"):
            inner += f"<p>{E(s['text'])}</p>"
        if s.get("table"):
            inner += f'<div style="margin-top:18px">{h_table(s["table"])}</div>'
        if s.get("bullets"):
            inner += "<ul>" + "".join(f"<li>{E(b)}</li>" for b in s["bullets"]) + "</ul>"
        if s.get("note"):
            inner += f'<p style="font-size:17px;color:#6B675F;margin-top:18px">{E(s["note"])}</p>'
        inner += "</div>"
        if s.get("image"):
            inner += h_img(s["image"], 830, 640, 1010, 360)
    elif t == "full":
        inner = (h_img(s["image"], 0, 0, 1920, 1080)
                 + f'<div style="position:absolute;left:80px;top:70px;background:rgba(246,243,238,.93);padding:26px 34px;width:640px">'
                   f'<div class="eyebrow">{E(s["eyebrow"])}</div><h2 style="font-size:40px">{E(s["title"])}</h2><p style="font-size:19px">{E(s["text"])}</p></div>')
    elif t == "grid":
        inner = f'<div class="bar"></div><div style="position:absolute;left:80px;top:104px;width:1700px"><div class="eyebrow">{E(s["eyebrow"])}</div><h2>{E(s["title"])}</h2></div>'
        for i, im in enumerate(s["images"]):
            inner += h_img(im, 80 + (i % 3) * 593, 280 + (i // 3) * 362, 573, 342)
    elif t == "materials":
        inner = (f'<div class="bar"></div><div style="position:absolute;left:80px;top:104px;width:620px"><div class="eyebrow">{E(s["eyebrow"])}</div><h2>{E(s["title"])}</h2>'
                 f'<div style="margin-top:18px">{h_table(s["table"])}</div><p style="font-size:16px;color:#6B675F;margin-top:14px">{E(s["note"])}</p></div>')
        inner += h_img(s["image"], 760, 104, 1080, 896)
    elif t == "table":
        inner = (f'<div class="bar"></div><div style="position:absolute;left:80px;top:104px;width:1760px"><div class="eyebrow">{E(s["eyebrow"])}</div><h2>{E(s["title"])}</h2>'
                 f'<div style="margin-top:24px">{h_table(s["table"])}</div><p style="font-size:17px;color:#6B675F;margin-top:16px">{E(s["note"])}</p></div>')
    elif t == "hero":
        dark = True
        inner = (h_img(s["image"], 0, 0, 1920, 1080)
                 + '<div style="position:absolute;inset:0;background:linear-gradient(0deg,rgba(15,30,21,.9) 0%,rgba(15,30,21,.3) 45%,rgba(15,30,21,0) 70%)"></div>'
                 + f'<div style="position:absolute;left:80px;bottom:100px;width:1300px"><div class="eyebrow" style="color:#9FD3B0">{E(s["eyebrow"])}</div>'
                   f'<h1 style="color:#fff">{E(s["title"])}</h1><p style="color:#D9E2DC;font-size:24px">{E(s["text"])}</p></div>')
    foot = (f'<div class="foot"><span><b>STAGE II KNITTING APPAREL / ZARA FASHION</b> &nbsp;|&nbsp; Ground-floor corporate office - interior design proposal</span>'
            f'<span>{n:02d} / {len(SPEC)}</span></div>')
    return f'<section class="slide{" dark" if dark else ""}">{inner}{foot}</section>'


html_doc = (f'<!doctype html><html><head><meta charset="utf-8"><title>Garment Office Design Pitch</title><style>{CSS}</style></head><body>'
            + "".join(h_slide(i + 1, s) for i, s in enumerate(SPEC)) + "</body></html>")
with open(os.path.join(OUT, "Garment_Office_Design_Pitch.html"), "w") as f:
    f.write(html_doc)

# ------------------------------------------------------------------------------------------ PPTX renderer (editable)
try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    PX = 6350                    # EMU per px for a 1920 x 1080 canvas on a 13.333 x 7.5 in slide
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)
    blank = prs.slide_layouts[6]
    GREEN, INK, GREY, LIGHT = RGBColor(0x26, 0x67, 0x3D), RGBColor(0x1E, 0x1E, 0x1C), RGBColor(0x5A, 0x57, 0x50), RGBColor(0xF6, 0xF3, 0xEE)

    def box(sl, x, y, w, h, text, size=20, bold=False, color=INK, font="Inter"):
        tb = sl.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
        tf = tb.text_frame
        tf.word_wrap = True
        lines = text if isinstance(text, list) else [text]
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = ln
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
            r.font.color.rgb = color
        return tb

    def pic(sl, spec_img, x, y, w, h):
        name, caption = spec_img
        if name.endswith((".jpg", ".png")):
            p = os.path.join(ASSET, name)
        else:
            p = os.path.join(ASSET, name + ".jpg")
        if not os.path.exists(p):
            return
        from PIL import Image
        im = Image.open(p)
        iw, ih = im.size
        # crop to fill (object-fit: cover)
        r_box, r_img = w / h, iw / ih
        shape = sl.shapes.add_picture(p, Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
        if r_img > r_box:
            c = (1 - r_box / r_img) / 2
            shape.crop_left = shape.crop_right = c
        else:
            c = (1 - r_img / r_box) / 2
            shape.crop_top = shape.crop_bottom = c
        if caption:
            box(sl, x + 10, y + h - 34, w - 20, 30, caption + (" (preview render)" if USED.get(name) == "preview" else ""), 11, False, RGBColor(0xFF, 0xFF, 0xFF))

    def bg(sl, color):
        f = sl.background.fill
        f.solid()
        f.fore_color.rgb = color

    def table(sl, rows, x, y, w, row_h=46, size=12):
        tbl = sl.shapes.add_table(len(rows), len(rows[0]), Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(row_h * len(rows) * PX)).table
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = tbl.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r else RGBColor(0xEE, 0xEA, 0xE2)
                cell.text = str(val)
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(size if r else size - 2)
                        run.font.bold = (r == 0) or (c == 0)
                        run.font.name = "Inter"
                        run.font.color.rgb = GREEN if (r == 0 or c == 0) else INK

    def panel(sl, x, y, w, h, color):
        shp = sl.shapes.add_shape(1, Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
        shp.fill.solid()
        shp.fill.fore_color.rgb = color
        shp.line.fill.background()
        return shp

    for n, s in enumerate(SPEC, 1):
        sl = prs.slides.add_slide(blank)
        bg(sl, LIGHT)
        t = s["type"]
        WHITE, MINT = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x9F, 0xD3, 0xB0)
        if t == "cover":
            pic(sl, s["image"], 0, 0, 1920, 1080)
            panel(sl, 0, 0, 1040, 1080, RGBColor(0x14, 0x24, 0x1B))
            panel(sl, 60, 60, 900, 110, WHITE)
            for i, (lg, wpx) in enumerate((("logo_zara_fashion_profile.png", 250), ("logo_stage2_ltd.png", 250), ("logo_stage2_unit1.png", 250))):
                sl.shapes.add_picture(os.path.join(ASSET, lg), Emu((80 + i * 290) * PX), Emu(78 * PX), width=Emu(wpx * PX))
            box(sl, 80, 300, 900, 40, s["eyebrow"].upper(), 13, True, MINT)
            box(sl, 80, 350, 900, 420, s["title"], 36, True, WHITE)
            box(sl, 80, 790, 900, 120, s["sub"], 15, False, RGBColor(0xD9, 0xE2, 0xDC))
            sl.shapes.add_picture(os.path.join(ASSET, "tagline_quality_compliance_commitment.png"), Emu(80 * PX), Emu(930 * PX), width=Emu(620 * PX))
        elif t == "hero":
            pic(sl, s["image"], 0, 0, 1920, 1080)
            panel(sl, 0, 600, 1920, 480, RGBColor(0x14, 0x24, 0x1B))
            box(sl, 80, 640, 1500, 40, s["eyebrow"].upper(), 13, True, MINT)
            box(sl, 80, 690, 1600, 170, s["title"], 32, True, WHITE)
            box(sl, 80, 870, 1500, 120, s["text"], 15, False, RGBColor(0xD9, 0xE2, 0xDC))
        elif t == "full":
            pic(sl, s["image"], 0, 0, 1920, 1080)
            panel(sl, 50, 50, 720, 330, LIGHT)
            box(sl, 80, 76, 660, 40, s["eyebrow"].upper(), 13, True, GREEN)
            box(sl, 80, 120, 660, 130, s["title"], 26, True, INK)
            box(sl, 80, 260, 660, 110, s["text"], 13, False, GREY)
        elif t == "room":
            box(sl, 80, 96, 580, 40, s["eyebrow"].upper(), 13, True, GREEN)
            box(sl, 80, 136, 580, 240, s["title"], 24, True)
            box(sl, 80, 400, 580, 620, [s["text"], ""] + ["- " + b for b in s["bullets"]], 14, False, GREY)
            ims = s["images"]
            lay = s.get("layout", "1+2")
            if lay == "1":
                pic(sl, ims[0], 700, 104, 1140, 896)
            elif lay == "2":
                pic(sl, ims[0], 700, 104, 1140, 440)
                pic(sl, ims[1], 700, 560, 1140, 440)
            else:
                pic(sl, ims[0], 700, 104, 1140, 560)
                pic(sl, ims[1], 700, 680, 562, 320)
                pic(sl, ims[2], 1278, 680, 562, 320)
        elif t == "cards":
            box(sl, 80, 96, 1500, 40, s["eyebrow"].upper(), 13, True, GREEN)
            box(sl, 80, 136, 1500, 120, s["title"], 28, True)
            for i, (ct, cd) in enumerate(s["cards"]):
                x, y = 80 + (i % 3) * 600, 320 + (i // 3) * 300
                panel(sl, x, y, 560, 270, WHITE)
                box(sl, x + 24, y + 22, 510, 50, (f"{i + 1:02d}   " if s.get("numbered") else "") + ct, 19, True, GREEN)
                box(sl, x + 24, y + 86, 510, 170, cd, 14, False, GREY)
            if s.get("image"):
                pic(sl, s["image"], 1280, 620, 560, 270)
        elif t == "plan":
            pic(sl, (s["plan"], ""), 80, 56, 680, 962)
            box(sl, 830, 96, 1010, 40, s["eyebrow"].upper(), 13, True, GREEN)
            box(sl, 830, 136, 1010, 130, s["title"], 26, True)
            y = 290
            if s.get("text"):
                box(sl, 830, y, 1010, 110, s["text"], 14, False, GREY)
                y += 120
            if s.get("table"):
                table(sl, s["table"], 830, y, 1010)
                y += 46 * len(s["table"]) + 24
            if s.get("bullets"):
                box(sl, 830, y, 1010, 320, ["- " + b for b in s["bullets"]], 14, False, GREY)
            if s.get("image"):
                pic(sl, s["image"], 830, 640, 1010, 360)
            if s.get("note"):
                box(sl, 830, 940, 1010, 70, s["note"], 11, False, GREY)
        elif t == "grid":
            box(sl, 80, 96, 1500, 40, s["eyebrow"].upper(), 13, True, GREEN)
            box(sl, 80, 136, 1700, 120, s["title"], 28, True)
            for i, im in enumerate(s["images"]):
                pic(sl, im, 80 + (i % 3) * 593, 280 + (i // 3) * 362, 573, 342)
        elif t in ("materials", "table"):
            w = 640 if t == "materials" else 1760
            box(sl, 80, 96, w, 40, s["eyebrow"].upper(), 13, True, GREEN)
            box(sl, 80, 136, w, 120, s["title"], 28, True)
            table(sl, s["table"], 80, 290, w, 50 if t == "table" else 46, 12 if t == "materials" else 13)
            box(sl, 80, 950, w, 60, s["note"], 11, False, GREY)
            if t == "materials":
                pic(sl, s["image"], 760, 104, 1080, 896)
        box(sl, 80, 1036, 1300, 30, "STAGE II KNITTING APPAREL / ZARA FASHION  |  Ground-floor corporate office - interior design proposal", 9, False, GREY)
        box(sl, 1740, 1036, 120, 30, f"{n:02d} / {len(SPEC)}", 9, False, GREY)
    prs.save(os.path.join(OUT, "Garment_Office_Design_Pitch.pptx"))
    print("pptx saved")
except Exception as e:  # keep the PDF path working even if python-pptx fails
    print("PPTX FAILED:", repr(e))

print("slides", len(SPEC), "| images:", {k: v for k, v in USED.items()})
